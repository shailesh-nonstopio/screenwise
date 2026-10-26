"""Resume upload validation, text extraction and structured parsing.

Parsing never invents details: anything that cannot be read with confidence is recorded as
unknown or uncertain, and extraction failures are reported to the user instead of guessed around.
"""

from __future__ import annotations

import io
import logging
import re
import zipfile
from dataclasses import dataclass
from datetime import date
from typing import Any

from pydantic import ValidationError

from app.config import settings
from app.schemas import ParsedResume
from app.services.llm_provider import LLMError, LLMProvider
from app.services.text_utils import clip

log = logging.getLogger(__name__)

PDF_MIME = "application/pdf"
DOCX_MIME = "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
MAX_TEXT_CHARS = 60_000
MAX_LINES = 300


class UploadError(ValueError):
    """A user-facing problem with an uploaded file."""


class ExtractionError(ValueError):
    """The file is valid but no usable text could be extracted."""


@dataclass
class ValidatedFile:
    kind: str  # pdf | docx
    mime: str
    safe_name: str


def safe_filename(name: str) -> str:
    base = (name or "resume").replace("\\", "/").split("/")[-1]
    base = re.sub(r"[^A-Za-z0-9._ -]", "_", base).strip(" .") or "resume"
    return base[:120]


def validate_upload(filename: str, content: bytes) -> ValidatedFile:
    max_bytes = settings().max_upload_bytes
    if not content:
        raise UploadError("The file is empty. Please choose your resume again.")
    if len(content) > max_bytes:
        raise UploadError(f"The file is larger than {settings().max_upload_mb:g} MB. Please upload a smaller file.")
    ext = filename.lower().rsplit(".", 1)[-1] if "." in filename else ""
    if content.startswith(b"%PDF-"):
        if ext not in ("pdf", ""):
            raise UploadError("The file's contents don't match its .%s extension." % ext)
        return ValidatedFile("pdf", PDF_MIME, safe_filename(filename))
    if content.startswith(b"PK\x03\x04"):
        if ext not in ("docx", ""):
            raise UploadError("Only PDF and DOCX files are accepted.")
        try:
            with zipfile.ZipFile(io.BytesIO(content)) as zf:
                infos = zf.infolist()
                if len(infos) > 2000 or sum(i.file_size for i in infos) > 80 * 1024 * 1024:
                    raise UploadError("This DOCX file is unusually large inside. Please re-save it and try again.")
                if "word/document.xml" not in zf.namelist():
                    raise UploadError("This doesn't look like a Word (.docx) document.")
        except zipfile.BadZipFile as exc:
            raise UploadError("This DOCX file appears to be damaged.") from exc
        return ValidatedFile("docx", DOCX_MIME, safe_filename(filename))
    if ext == "doc":
        raise UploadError("Older .doc files aren't supported. Please save as PDF or DOCX.")
    raise UploadError("Only PDF and DOCX files are accepted.")


def extract_text(content: bytes, kind: str) -> str:
    if kind == "pdf":
        text = _extract_pdf(content)
    elif kind == "docx":
        text = _extract_docx(content)
    else:  # pragma: no cover - guarded by validate_upload
        raise ExtractionError("Unsupported file type.")
    text = text.replace("\x00", "")
    text = re.sub(r"[ \t]+", " ", text)
    text = re.sub(r"\n{3,}", "\n\n", text).strip()
    if len(re.sub(r"\W", "", text)) < 40:
        raise ExtractionError(
            "We couldn't find readable text in this file. It may be a scanned image. "
            "Please upload a PDF or DOCX with selectable text."
        )
    return text[:MAX_TEXT_CHARS]


def _extract_pdf(content: bytes) -> str:
    from pypdf import PdfReader

    try:
        reader = PdfReader(io.BytesIO(content))
        if reader.is_encrypted:
            try:
                if not reader.decrypt(""):
                    raise ExtractionError("This PDF is password-protected. Please upload an unprotected copy.")
            except Exception as exc:
                raise ExtractionError("This PDF is password-protected. Please upload an unprotected copy.") from exc
        if len(reader.pages) > 30:
            raise ExtractionError("This PDF has more than 30 pages. Please upload your resume only.")
        parts = [(page.extract_text() or "") for page in reader.pages]
    except ExtractionError:
        raise
    except Exception as exc:  # pypdf raises many error types on malformed input  # noqa: BLE001
        raise ExtractionError("We couldn't read this PDF. It may be damaged; please re-save and upload it again.") from exc
    return "\n".join(parts)


def _extract_docx(content: bytes) -> str:
    import docx

    try:
        document = docx.Document(io.BytesIO(content))
    except Exception as exc:  # python-docx raises a variety of errors on bad input
        raise ExtractionError("We couldn't read this DOCX file. Please re-save it and try again.") from exc
    parts = [p.text for p in document.paragraphs]
    for table in document.tables:
        for row in table.rows:
            parts.append(" | ".join(cell.text.strip() for cell in row.cells if cell.text.strip()))
    return "\n".join(parts)


def number_lines(text: str) -> list[dict[str, str]]:
    lines = []
    for raw in text.splitlines():
        t = re.sub(r"^[\s•●▪\-*·]+", "", raw).strip()
        if len(t) < 2:
            continue
        lines.append({"ref": f"R-{len(lines) + 1}", "text": clip(t, 400)})
        if len(lines) >= MAX_LINES:
            break
    return lines


# --------------------------------------------------------------------------- heuristic parser

_MONTHS = {m: i + 1 for i, m in enumerate(["jan", "feb", "mar", "apr", "may", "jun", "jul", "aug", "sep", "oct", "nov", "dec"])}
_DATE = r"(?:(?P<{p}q>early|mid|late)\s+)?(?:(?P<{p}m>jan|feb|mar|apr|may|jun|jul|aug|sep|sept|oct|nov|dec)[a-z]*\.?\s+)?(?P<{p}y>(?:19|20)\d{{2}})"
_RANGE = re.compile(
    _DATE.format(p="s")
    + r"\s*(?:-|–|—|to|until)\s*(?:(?P<present>present|current|now|today)|"
    + _DATE.format(p="e")
    + r")",
    re.IGNORECASE,
)
_EDU = re.compile(r"\b(b\.?\s?e\b|b\.?\s?tech|b\.?\s?sc|b\.?a\b|bachelor|master|m\.?\s?sc|m\.?\s?tech|mba|ph\.?d|university|college|institute|diploma)", re.I)
_ACHIEVE = re.compile(r"\b(\d+%|\d+x|reduced|increased|improved|cut|grew|saved|launched|won|award|delivered)\b", re.I)
_PROJECT = re.compile(r"\b(project|built|developed|designed|created|implemented)\b", re.I)

KNOWN_SKILLS = [
    "Python", "Java", "JavaScript", "TypeScript", "Go", "Rust", "C#", "C++", "Ruby", "PHP", "Kotlin", "Swift", "Scala",
    "SQL", "PostgreSQL", "MySQL", "SQLite", "MongoDB", "Redis", "Elasticsearch", "Kafka", "RabbitMQ", "Celery",
    "FastAPI", "Django", "Flask", "Spring", "Node.js", "React", "Vue", "Angular", "Next.js", "GraphQL", "REST",
    "Docker", "Kubernetes", "Terraform", "AWS", "GCP", "Azure", "Linux", "Git", "CI/CD", "Grafana", "Prometheus",
    "pandas", "NumPy", "scikit-learn", "PyTorch", "TensorFlow", "Spark", "Airflow", "dbt", "Snowflake", "BigQuery",
    "Tableau", "Looker", "Power BI", "Excel", "Statistics", "A/B testing", "Machine Learning",
    "Salesforce", "HubSpot", "Zendesk", "Jira", "Figma", "Agile", "Scrum",
]


def _month(m: str | None) -> int | None:
    return _MONTHS.get(m[:3].lower()) if m else None


def _bounds(qual: str | None, month: int | None, year: int, *, start: bool) -> tuple[date, date, bool]:
    """Earliest and latest plausible date for a resume date expression."""
    if month:
        d = date(year, month, 1)
        return d, d, True
    if qual:
        q = qual.lower()
        lo, hi = {"early": (1, 4), "mid": (5, 8), "late": (9, 12)}[q]
        return date(year, lo, 1), date(year, hi, 1), False
    return date(year, 1, 1), date(year, 12, 1), False


def _months_between(a: date, b: date) -> int:
    return max(0, (b.year - a.year) * 12 + (b.month - a.month))


def _merge(intervals: list[tuple[date, date]]) -> int:
    if not intervals:
        return 0
    intervals.sort()
    total = 0
    cur_s, cur_e = intervals[0]
    for s, e in intervals[1:]:
        if s <= cur_e:
            cur_e = max(cur_e, e)
        else:
            total += _months_between(cur_s, cur_e)
            cur_s, cur_e = s, e
    total += _months_between(cur_s, cur_e)
    return total


def heuristic_parse(lines: list[dict[str, str]], extra_skills: list[str] | None = None, today: date | None = None) -> dict[str, Any]:
    today = today or date.today()
    roles: list[dict[str, Any]] = []
    lo_intervals: list[tuple[date, date]] = []
    hi_intervals: list[tuple[date, date]] = []
    education, achievements, projects = [], [], []
    all_certain = True

    for ln in lines:
        text = ln["text"]
        m = _RANGE.search(text)
        if m:
            sy = int(m.group("sy"))
            s_lo, s_hi, s_ok = _bounds(m.group("sq"), _month(m.group("sm")), sy, start=True)
            if m.group("present"):
                e_lo = e_hi = date(today.year, today.month, 1)
                e_ok, end_label = True, "present"
            else:
                ey = int(m.group("ey"))
                e_lo, e_hi, e_ok = _bounds(m.group("eq"), _month(m.group("em")), ey, start=False)
                end_label = m.group(0).split(m.group("sy"), 1)[1].lstrip(" -–—to").strip()
            start_label = text[m.start() : m.start() + len(m.group(0))].split("-")[0].split("–")[0].split("—")[0].strip()
            head = text[: m.start()].strip(" ,|·•-–—()")
            head = re.sub(r"\s*\(\s*$", "", head)
            parts = [p.strip() for p in re.split(r"\s+[—–-]\s+|\s+at\s+|\s*\|\s*|,\s+", head) if p.strip()]
            title = parts[0] if parts else ""
            org = parts[1] if len(parts) > 1 else None
            notes = []
            if not s_ok:
                notes.append(f"Start date is approximate (“{start_label}”)")
            if not e_ok:
                notes.append("End date is approximate")
            if not title:
                notes.append("Job title not found on this line")
            roles.append(
                {
                    "title": clip(title, 120) or "Unknown title",
                    "organization": clip(org, 120) if org else None,
                    "start": start_label,
                    "end": end_label,
                    "start_certain": s_ok,
                    "end_certain": e_ok,
                    "ref": ln["ref"],
                    "uncertain_note": "; ".join(notes),
                }
            )
            all_certain = all_certain and s_ok and e_ok
            if s_lo <= e_hi:
                # Minimum: latest plausible start to earliest plausible end; maximum: the reverse.
                if s_hi <= e_lo:
                    lo_intervals.append((s_hi, e_lo))
                hi_intervals.append((s_lo, e_hi))
            continue
        if _EDU.search(text):
            year = re.search(r"\b(19|20)\d{2}\b", text)
            education.append({"text": text, "ref": ln["ref"], "year": year.group(0) if year else None})
        if _ACHIEVE.search(text):
            achievements.append({"text": text, "ref": ln["ref"]})
        elif _PROJECT.search(text):
            projects.append({"text": text, "ref": ln["ref"]})

    full = "\n".join(ln["text"] for ln in lines)
    skills: list[str] = []
    for skill in KNOWN_SKILLS + list(extra_skills or []):
        pattern = r"(?<![A-Za-z0-9])" + re.escape(skill) + r"(?![A-Za-z0-9])"
        if re.search(pattern, full, re.I) and skill.lower() not in {s.lower() for s in skills}:
            skills.append(skill)

    exp = None
    if hi_intervals:
        lo = round(_merge(lo_intervals) / 12, 1)
        hi = round(_merge(hi_intervals) / 12, 1)
        exp = {
            "min": min(lo, hi),
            "max": hi,
            "certain": all_certain and abs(hi - lo) < 0.2,
            "note": "Total of listed roles, overlaps merged. Relevance to this job is not judged here.",
        }

    unknowns = []
    if not roles:
        unknowns.append("Work history dates (no date ranges found)")
    if roles and not all_certain:
        unknowns.append("Exact dates for some roles")
    if not education:
        unknowns.append("Education")
    elif not any(e["year"] for e in education):
        unknowns.append("Graduation year")
    if not skills:
        unknowns.append("Skills (none recognised)")

    return {
        "roles": roles[:30],
        "skills": skills[:80],
        "education": education[:15],
        "achievements": achievements[:30],
        "projects": projects[:30],
        "experience_years": exp,
        "unknowns": unknowns,
    }


def parse_resume(lines: list[dict[str, str]], provider: LLMProvider, extra_skills: list[str] | None = None) -> dict[str, Any]:
    """Return a validated structured resume. Falls back to the heuristic parser on any LLM problem."""
    valid_refs = {ln["ref"] for ln in lines}
    heuristic = heuristic_parse(lines, extra_skills)
    if provider.is_mock:
        result = ParsedResume.model_validate(heuristic).model_dump()
        result["source"] = "heuristic"
        return result
    try:
        raw = provider.parse_resume(lines)
        parsed = ParsedResume.model_validate(raw).model_dump()
    except (LLMError, ValidationError) as exc:
        log.warning("resume llm parse failed; using heuristic", extra={"safe": {"error": type(exc).__name__}})
        result = ParsedResume.model_validate(heuristic).model_dump()
        result["source"] = "heuristic"
        result["unknowns"].append("AI parsing unavailable; basic parser used")
        return result
    # Drop references that point at lines that do not exist.
    for key in ("roles", "education", "achievements", "projects"):
        for item in parsed[key]:
            if item.get("ref") not in valid_refs:
                item["ref"] = None
    # Never let the model drop recognised skills.
    for s in heuristic["skills"]:
        if s.lower() not in {x.lower() for x in parsed["skills"]}:
            parsed["skills"].append(s)
    parsed["source"] = "llm"
    return parsed
